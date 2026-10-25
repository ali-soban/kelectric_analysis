"""KE net-metering bill pipeline: fetch -> import -> extract -> workbook -> dashboard.

Double-click "Run KE Update.bat". The original "KE NM BILL.xlsx" is only ever read.
"""
import datetime as dt
import json
import math
import os
import re
import sys
import time
import webbrowser
from pathlib import Path

import openpyxl
import pdfplumber
from openpyxl.styles import Alignment, Font, PatternFill

KE_URL = "https://staging.ke.com.pk:24555/"
BILLS = Path(os.environ.get("KE_BILLS_DIR") or Path.home() / "Desktop" / "ke_bills")
ORIGINAL = BILLS / "KE NM BILL.xlsx"
AUTO = BILLS / "KE Bills Auto.xlsx"
DASH = BILLS / "KE Dashboard.html"
TEMPLATE = Path(__file__).with_name("dashboard_template.html")
DOWNLOADS = Path.home() / "Downloads"
FIRST_SHEET_BILL = (2023, 10)  # row 3 of the original sheet is the Oct-23 bill
MONTHS = "Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split()


class BillError(Exception):
    pass


def num(s):
    return float(s.replace(",", ""))


def half_up(x):
    return int(x + 0.5) if x >= 0 else -int(-x + 0.5)


def pdf_text(path):
    with pdfplumber.open(path) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def bill_month(text):
    """'2609' from the header line ': Sep-26 : 07-Sep-2026' (new format only)."""
    m = re.search(r"^: (\w{3})-(\d{2}) : \d", text, re.M)
    return f"{m[2]}{MONTHS.index(m[1]) + 1:02d}" if m and m[1] in MONTHS else None


def yymm_to_date(yymm):
    return dt.date(2000 + int(yymm[:2]), int(yymm[2:]), 1)


def label(yymm):  # '2609' -> 'Sep-26'
    return f"{MONTHS[int(yymm[2:]) - 1]}-{yymm[:2]}"


# ---------------------------------------------------------------- extraction

def parse_bill(path):
    """Extract one new-format (Aug-25 onward) bill. Raises BillError rather than guessing."""
    t = pdf_text(path)
    ym = bill_month(t)
    if not ym or "Variable Off Peak" not in t:
        raise BillError("not a new-format KE bill")

    def find(pattern, what):
        m = re.search(pattern, t, re.M)
        if not m:
            raise BillError(f"could not find {what}")
        return m

    def meter(name):  # "<name> <previous> <current> <units>"
        return float(find(rf"^{name} [\d.]+ [\d.]+ ([\d.]+)", name)[1])

    def money(name):
        return num(find(rf"^{name} (-?[\d,]+(?:\.\d+)?)", name)[1])

    off, peak = meter("Energy - Off Peak"), meter("Energy - Peak")
    exported = meter("Import Energy Off Peak") + meter("Import Energy Peak")

    # One line per tariff in force; a mid-period change adds a "(New)" line.
    var_off = re.findall(r"^Variable Off Peak(?: \(New\))? ([\d.]+) ([\d.]+) ", t, re.M)
    var_peak = re.findall(r"^Variable Peak(?: \(New\))? ([\d.]+) ([\d.]+) ", t, re.M)
    if not var_off or not var_peak:
        raise BillError("could not find tariff lines")
    if abs(sum(float(u) for u, _ in var_off) - off) > 0.05 * len(var_off):
        raise BillError("off-peak billed units do not match the meter reading")

    billed = money("Current Bill Amount")
    benefit = money("Total Net Metering Benefit")
    opening = money("Opening Balance")
    prev = money(r"Previous (?:Recoverable )?Dues")
    net = prev + billed - benefit  # what this bill leaves: + to pay, - credit

    outstanding = re.search(r"^Outstanding Balance (-?[\d,]+\.\d+)", t, re.M)
    printed = re.search(r"^Amount Payable within Due Date ([\d,]+)", t, re.M)
    if outstanding:
        payable, credit = 0, half_up(-num(outstanding[1]))
        if abs(num(outstanding[1]) - net) > 1:
            raise BillError(f"outstanding balance {outstanding[1]} != computed {net:.2f}")
    elif printed:
        payable, credit = int(num(printed[1])), 0
        if abs(payable - math.ceil(round(net, 2))) > 1:
            raise BillError(f"printed payable {printed[1]} != computed {net:.2f}")
    elif net < -1:
        payable, credit = 0, half_up(-net)
    elif net <= 1:
        # Downloaded after paying: the payment already sits in Previous Dues, so
        # rebuild what was due from the credit (if any) carried into this bill.
        payable, credit = math.ceil(round(min(0.0, opening) + billed - benefit, 2)), 0
    else:
        payable, credit = math.ceil(round(net, 2)), 0

    ppp = re.search(r"NAPP Off-Peak ([\d.,]+) ([\d.]+) ", t)
    head = find(r"^Description (\w{3}-\d{2}) (\w{3}-\d{2}) (\w{3}-\d{2})$", "carry-forward history")
    tot = find(r"^Total ([\d,.]+) ([\d,.]+) ([\d,.]+)$", "carry-forward totals")
    if label(ym) not in head.groups():
        raise BillError("carry-forward history does not include this month")
    cf = 0.0 if ppp else num(tot[head.groups().index(label(ym)) + 1])

    reading = find(r"Reading Date: (\d{2}-\w{3}-\d{2})", "reading date")[1]
    return {
        "breakdown": charge_breakdown(t, billed),
        "bill": ym, "source": "PDF",
        "reading": dt.datetime.strptime(reading, "%d-%b-%y").date(),
        "E": exported, "F": off, "G": peak,
        "I": float(var_off[-1][1]), "J": float(var_peak[-1][1]),
        "L": cf, "M": payable, "N": billed, "Q": credit, "R": benefit,
        "S": num(ppp[1]) if ppp else None, "T": float(ppp[2]) if ppp else None,
    }


CHARGE_GROUPS = [("Variable Off Peak", "Off-peak energy"), ("Variable Peak", "Peak energy"),
                 ("Fixed Charges", "Fixed charges"), ("FCA", "Fuel adjustment (FCA)"),
                 ("Uniform Quarterly", "Quarterly adjustment"), ("Additional Surcharge", "PHL surcharge"),
                 ("Electricity Duty", "Electricity duty"), ("Sales Tax", "Sales tax")]


def charge_breakdown(t, billed):
    """Bill charges grouped by type; None unless they add up to the Current Bill Amount."""
    sec = re.search(r"Units Rate Amount\n(.*?)\nElectricity Charges for current month", t, re.S)
    if not sec:
        return None
    groups = {}
    for line in sec[1].splitlines():
        amt = re.search(r"(-?[\d,]+\.\d{2})$", line)
        if not amt or line.startswith(("Electricity Charges ", "Taxes and Duties")):
            continue  # subtotals
        name = next((g for prefix, g in CHARGE_GROUPS if line.startswith(prefix)),
                    re.sub(r"[\s\d.,:-]+$", "", line[:amt.start()]).strip())
        groups[name] = round(groups.get(name, 0) + num(amt[1]), 2)
    return groups if abs(sum(groups.values()) - billed) <= 1 else None


def first_number(v):
    m = re.search(r"\d+(?:\.\d+)?", str(v or ""))
    return float(m[0]) if m else None


def read_original():
    """Hand-entered rows of the original sheet, keyed by bill (row 3 = 2310). Values only."""
    ws = openpyxl.load_workbook(ORIGINAL, data_only=True, read_only=True).worksheets[0]
    rows, (y, m) = {}, FIRST_SHEET_BILL
    for r in ws.iter_rows(min_row=3, max_col=20, values_only=True):
        if r[13] is None:  # N (Billed) empty -> no bill yet
            break
        v = dict(zip("ABCDEFGHIJKLMNOPQRST", r))
        rows[f"{y % 100:02d}{m:02d}"] = {
            "C": v["C"], "E": v["E"], "F": v["F"], "G": v["G"], "I": v["I"], "J": v["J"],
            "L": v["L"], "M": v["M"], "N": v["N"], "Q": v["Q"], "R": v["R"],
            "S": first_number(v["S"]), "T": first_number(v["T"]),
        }
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return rows


def read_solar():
    """Solar production you typed into the generated workbook (kept across runs)."""
    if not AUTO.exists():
        return {}
    while True:  # wait until Excel lets go, so an unsaved entry gets saved before we read
        try:
            open(AUTO, "a").close()
            break
        except PermissionError:
            input(f"  {AUTO.name} is open in Excel. Save and close it, then press Enter: ")
    ws = openpyxl.load_workbook(AUTO, read_only=True, data_only=True)["Bills"]
    solar = {}
    for bill, _, c in ws.iter_rows(min_row=2, max_col=3, values_only=True):
        if not (bill and str(bill).isdigit()) or c is None:
            continue
        if isinstance(c, (int, float)):
            solar[str(bill)] = c
        else:
            print(f"  ! Solar Production for {label(str(bill))} is not a number ({c!r}), ignored")
    return solar


# ---------------------------------------------------------------- fetch / import

def account_number():
    for pdf in sorted(BILLS.glob("[0-9][0-9][0-9][0-9].pdf"), reverse=True):
        m = re.search(r"Account No\. (\d{10,})", pdf_text(pdf))
        if m:
            return m[1]
    return None


def import_downloads(account):
    """Move KE bills for this account from Downloads into ke_bills as YYMM.pdf."""
    cutoff = time.time() - 120 * 86400
    for f in DOWNLOADS.glob("*.pdf"):
        if f.stat().st_mtime < cutoff:
            continue
        try:
            text = pdf_text(f)
        except Exception:
            continue
        if "Net Metering Invoice" not in text or (account and account not in text):
            continue
        ym = bill_month(text)
        if not ym:
            print(f"  ? {f.name}: KE bill but bill month not found, left in Downloads")
            continue
        dest = BILLS / f"{ym}.pdf"
        if dest.exists():
            if dest.read_bytes() == f.read_bytes():
                f.unlink()
                print(f"  {f.name}: already saved as {dest.name}, removed duplicate")
            else:
                print(f"  ! {f.name}: {dest.name} already exists with different content, left in Downloads")
            continue
        f.replace(dest)
        print(f"  {f.name} -> {dest.name}")


def fetch_latest(account):
    today = dt.date.today()
    current = f"{today.year % 100:02d}{today.month:02d}"
    import_downloads(account)
    if (BILLS / f"{current}.pdf").exists():
        print(f"  {current}.pdf already here.")
        return
    print(f"\n  {label(current)} bill not found. Opening the KE bill page...")
    print("  Type the characters, download the PDF, then come back here.")
    webbrowser.open(KE_URL)
    input("  Press Enter when downloaded (or just Enter to skip): ")
    import_downloads(account)


# ---------------------------------------------------------------- outputs

HEADERS = [
    ("Bill", 7), ("Bill Month", 10), ("Solar Production (kWh)", 12), ("Solar Used (kWh)", 10),
    ("Solar Exported (kWh)", 10), ("Off-Peak Used (kWh)", 10), ("On-Peak Used (kWh)", 10),
    ("Net Off-Peak Saved (kWh)", 11), ("Off-Peak Rate", 9), ("On-Peak Rate", 9),
    ("Avg Price per Unit", 9), ("C/F Units (kWh)", 9), ("Bill Payable (PKR)", 12),
    ("Billed (PKR)", 12), ("Bill without Solar (PKR)", 13), ("Amount Saved (PKR)", 12),
    ("Credit Balance (PKR)", 12), ("NM Benefit (PKR)", 12), ("PPP Units (kWh)", 9),
    ("PPP Rate", 8), ("Reading Date", 11), ("Source", 8),
]
CHECK_FIELDS = {"E": "Solar Exported", "F": "Off-Peak Used", "G": "On-Peak Used",
                "I": "Off-Peak Rate", "J": "On-Peak Rate", "L": "C/F Units",
                "M": "Bill Payable", "N": "Billed", "Q": "Credit Balance",
                "R": "NM Benefit", "S": "PPP Units", "T": "PPP Rate"}


def derive(row):
    """Same maths as the sheet formulas, for the dashboard."""
    c, e, f, g, n, m = (row.get(k) for k in "CEFGNM")
    k = n / (f + g) if f + g else None
    d = c - e if c is not None else None
    o = n + d * k if d is not None and k is not None else None
    return {**row, "D": d, "H": e - f, "K": k, "O": o, "P": o - m if o is not None else None}


def write_workbook(rows, original):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bills"
    bold, wrap = Font(bold=True), Alignment(wrap_text=True, vertical="top")
    todo = PatternFill("solid", fgColor="FFF2CC")
    for i, (name, width) in enumerate(HEADERS, 1):
        cell = ws.cell(1, i, name)
        cell.font, cell.alignment = bold, wrap
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "C2"

    for r, row in enumerate(rows, 2):
        vals = {
            "A": row["bill"], "B": yymm_to_date(row["bill"]), "C": row.get("C"),
            "D": f'=IF(C{r}="","",C{r}-E{r})', "E": row["E"], "F": row["F"], "G": row["G"],
            "H": f"=E{r}-F{r}", "I": row["I"], "J": row["J"], "K": f"=N{r}/(F{r}+G{r})",
            "L": row["L"], "M": row["M"], "N": row["N"],
            "O": f'=IF(C{r}="","",N{r}+D{r}*K{r})', "P": f'=IF(O{r}="","",O{r}-M{r})',
            "Q": row["Q"], "R": row["R"], "S": row["S"], "T": row["T"],
            "U": row.get("reading"), "V": row["source"],
        }
        for col, v in vals.items():
            ws[f"{col}{r}"] = v
        ws[f"B{r}"].number_format = "mmm-yy"
        ws[f"U{r}"].number_format = "dd-mmm-yy"
        for col in "CDEFGHLS":
            ws[f"{col}{r}"].number_format = "#,##0"
        for col in "IJKT":
            ws[f"{col}{r}"].number_format = "0.00"
        for col in "MNOPQR":
            ws[f"{col}{r}"].number_format = "#,##0"
        if row.get("C") is None:
            ws[f"C{r}"].fill = todo
    total = len(rows) + 2
    ws[f"A{total}"], ws[f"A{total}"].font = "Total", bold
    ws[f"P{total}"] = f"=SUM(P2:P{total - 1})"
    ws[f"P{total}"].font, ws[f"P{total}"].number_format = bold, "#,##0"
    ws[f"A{total + 2}"] = ("Type only in Solar Production: the previous calendar month's total "
                           "(bill Oct-26 gets September). Everything else is rebuilt from the PDFs on each run.")

    chk = wb.create_sheet("Check")
    chk.append(["Bill", "Field", "Your sheet", "PDF", "Difference"])
    for cell in chk[1]:
        cell.font = bold
    diffs = 0
    for row in rows:
        mine = original.get(row["bill"])
        if row["source"] != "PDF" or not mine:
            continue
        for key, name in CHECK_FIELDS.items():
            a, b = mine.get(key), row.get(key)
            tol = 0.005 if key in "IJT" else 1
            if (a or 0) != (b or 0) and abs((a or 0) - (b or 0)) > tol:
                chk.append([label(row["bill"]), name, a, round(b, 2) if b is not None else None,
                            round((b or 0) - (a or 0), 2)])
                diffs += 1
    for col, w in zip("ABCDE", (8, 16, 12, 12, 12)):
        chk.column_dimensions[col].width = w

    wb.save(AUTO)
    return diffs


def write_dashboard(rows):
    data = []
    for row in rows:
        d = derive(row)
        data.append({k: (v.isoformat() if isinstance(v, dt.date) else v) for k, v in d.items()})
    html = TEMPLATE.read_text(encoding="utf-8")
    payload = json.dumps({"generated": dt.date.today().isoformat(), "rows": data})
    DASH.write_text(html.replace("__DATA__", payload.replace("</", "<\\/")), encoding="utf-8")


# ---------------------------------------------------------------- main

def main():
    print("KE bill update\n\n1. Fetch")
    account = account_number()
    fetch_latest(account)

    print("\n2. Extract")
    original, solar = read_original(), read_solar()
    rows, failed = {}, []
    for bill, v in original.items():
        rows[bill] = {"bill": bill, "source": "manual", **v}
    for pdf in sorted(BILLS.glob("[0-9][0-9][0-9][0-9].pdf")):
        try:
            p = parse_bill(pdf)
        except BillError as e:
            if pdf.stem >= "2508":  # older layout is expected to be skipped
                failed.append(f"{pdf.name}: {e}")
            continue
        if p["bill"] != pdf.stem:
            failed.append(f"{pdf.name}: PDF says it is the {label(p['bill'])} bill")
            continue
        rows[pdf.stem] = p
    for bill, row in rows.items():
        row["C"] = solar.get(bill, original.get(bill, {}).get("C"))
    rows = [rows[b] for b in sorted(rows)]
    n_pdf = sum(r["source"] == "PDF" for r in rows)
    print(f"  {n_pdf} bills read from PDF, {len(rows) - n_pdf} from your original sheet")
    for f in failed:
        print(f"  ! {f}")

    print("\n3. Workbook")
    diffs = write_workbook(rows, original)
    print(f"  Saved {AUTO.name} ({diffs} differences vs your sheet listed on the Check tab)")
    missing = [label(r["bill"]) for r in rows if r.get("C") is None]
    if missing:
        print(f"  Solar Production still needed for: {', '.join(missing)}")

    print("\n4. Dashboard")
    write_dashboard(rows)
    print(f"  Saved {DASH.name}")
    os.startfile(DASH)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
